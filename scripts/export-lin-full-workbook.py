#!/usr/bin/env python3
"""Build the full reviewed EC600 LIN workbook. Requires openpyxl (3.1+)."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs/LIN-Table-Layout-EC600-Full.xlsx'
catalogue = json.loads((ROOT / 'lin/ec600-lin-scheme.json').read_text())
entries = {e['id']: e for e in catalogue['entries']}
wb = Workbook()
wb.remove(wb.active)
NAVY, TEAL, AMBER, GREY = '17324D', 'DDF2EF', 'FFF0CC', 'EDF0F4'


def sheet(name, title, subtitle, headers, rows, widths):
    ws = wb.create_sheet(name)
    end = get_column_letter(len(headers))
    for n, text in [(1, title), (2, subtitle)]:
        ws.merge_cells(f'A{n}:{end}{n}')
        cell = ws.cell(n, 1, text)
        cell.fill = PatternFill('solid', fgColor=NAVY if n == 1 else TEAL)
        cell.font = Font(name='Calibri', size=18 if n == 1 else 11,
                         bold=n == 1, color='FFFFFF' if n == 1 else NAVY)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 44
    for col, text in enumerate(headers, 1):
        c = ws.cell(4, col, text)
        c.font = Font(name='Calibri', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 32
    for r, row in enumerate(rows, 5):
        for col, value in enumerate(row, 1):
            c = ws.cell(r, col, value)
            c.font = Font(name='Calibri', size=11, color=NAVY)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(value, str) and any(term in value.lower() for term in ('unknown', 'not established', 'not defined', 'suspect', 'conflict', 'disabled')):
                c.fill = PatternFill('solid', fgColor=AMBER)
        lines = max(max(len(str(v)) // max(int(widths[i]) - 3, 1) + 1,
                        str(v).count('\n') + 1) for i, v in enumerate(row))
        ws.row_dimensions[r].height = max(36, min(150, 16 * (lines + 1)))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    table = Table(displayName=name.replace(' ', '') + 'Table', ref=f'A4:{end}{ws.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'D5' if len(headers) > 5 else 'A5'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{end}{ws.max_row}'
    ws.oddFooter.center.text = 'EC600 LIN reference | Page &P of &N'
    return ws

# Structured field definitions below transcribe the reviewed catalogue. The byte
# map always retains the full catalogue text, including qualifications.
# Fail closed if that catalogue changes: the curated fields must be reviewed too.
import hashlib
CATALOGUE_HASH = '62def9f367c8ac742cd4054ee612e1ce5ed8b91f234e6f3746eaeadfd1e7817a'
assert hashlib.sha256((ROOT / 'lin/ec600-lin-scheme.json').read_bytes()).hexdigest() == CATALOGUE_HASH, 'Catalogue changed: review curated signal definitions before updating its hash.'

fields = {}

def define(ids, device, byte, *items):
    """Each item is (mask, signal, values/conversion[, evidence])."""
    for ident in ids:
        fields[ident, device, byte] = list(items)


def flag(mask, name):
    return (mask, name, '0=flag clear; 1=flag set. Interpret in the selected device/context.')


def unknown(mask, note='Meaning not established by the reviewed PSU code.'):
    return (mask, 'Unknown / undecoded bits', note, 'Unknown')

# Air conditioning.
for ident, byte in [(8, 3), (23, 5)]:
    define([ident], 'Truma aircon', byte,
           (0x0F, 'Operating mode', '0=off; 4=fan; 5=cooling; 6=heating; 7=automatic.'), unknown(0xF0))
for ident in (8, 23):
    define([ident], 'Dometic aircon', 0,
           flag(0x80, 'Aircon on / mode component'), flag(0x40, 'Light on'),
           (0x03, 'Mode component', 'Part of composite mode; see Combined mode row.'),
           *([flag(0x04, 'Fan automatic'), unknown(0x38)] if ident == 8 else [unknown(0x3C)]))
    define([ident], 'Dometic aircon', 1,
           (0xF0, 'Target temperature', 'degrees C = extracted field + 16.'),
           (0x0C, 'Fan setting', 'Raw code 0..3; enumeration not established.'),
           (0x03, 'Mode component', 'Part of composite mode; see Combined mode row.'))
    define([ident], 'Dometic aircon', 4,
           (0xF0, 'iFeel temperature', 'degrees C = extracted field + 16.'),
           unknown(0x0F, 'Low nibble preserved in control; meaning not established.'))
define([8], 'Dometic aircon', 7, (0x04, 'Synchronisation flag', '0x04 set for sync frame; cleared after synchronising.'), unknown(0xFB))
define([23], 'Dometic aircon', 7, flag(0x01, 'Remote/settings change'), unknown(0xFE))

# Refrigerator: ID 11 echo reuses the ID 12 fields below.
for ident in (11, 12):
    define([ident], 'Dometic fridge', 0,
           (0x07, 'Operating mode', '0=off; 1=auto; 3=gas; 5=12 V; 7=230 V. Other codes not established.'),
           (0x01, 'On flag (overlaps mode)', '0=off; 1=on. Part of mode, not an independent setting.'), unknown(0xF8))
define([11], 'Dometic fridge', 5,
       (0x10, 'Automatic mode available', 'Set to 1 in both emitted values (0x10, 0x14).'),
       (0x04, 'Frame heater fitted', '0=not fitted; 1=fitted.'),
       (0xEB, 'Other control bits', 'Written as zero in this control path.'))
define([12], 'Dometic fridge', 4,
       (0x07, 'FridgeState', 'Raw 0..7; enumeration not established.'),
       flag(0x08, 'Manual mode'), flag(0x10, 'Door open'), unknown(0xE0))

# Alde generations and clock frames must remain separate.
define([26], 'Alde 3020+', 3,
       (0x3F, 'Room target', 'Sent field=(target degrees C - 5)*2, or 0 when off. Non-off inverse: degrees C=field/2+5.'),
       flag(0x40, 'Gas enabled'), (0x80, 'Not set by sender', 'Written as zero by this routine.'))
define([26], 'Alde 3020+', 4,
       (0xC0, 'Electric selection', 'Selection code 0..3; physical power units not established here.'),
       (0x3F, 'HeatingNow contribution', 'Code can add (HeatingNow - 5)*2 in low bits; conditional source behaviour, not a universal temperature field.'))
define([26], 'Alde 3020+', 5,
       flag(0x01, 'Alde on'), (0x18, 'Hot-water selection', '0=off; 1=normal; 2=boost.'),
       (0xE6, 'Other sender bits', 'Not set by this routine.'))
legacy_settings = [flag(0x01, 'Heater on'), flag(0x02, 'Gas'),
                   (0x0C, 'Electric power', '0..3 kW.'),
                   (0x30, 'Water setting', '0=off; 2=normal; 3=boost (returns to 2 after timer).'), unknown(0xC0)]
define([26, 27], 'Alde / EC645 legacy', 0, *legacy_settings)
define([26], 'Alde / EC645 clock', 3,
       (0x80, 'Clock marker', '1=clock frame. Ordinary legacy control has B3=0x40.'),
       (0x7F, 'Day field', 'Sent field=Day-1. Day numbering convention is not established here.'))
define([27], 'Alde 3020+', 3,
       (0x3F, 'Temperature setting', 'degrees C=(extracted field + 10)/2.'),
       flag(0x40, 'Gas'), (0x80, 'Energy priority', 'Raw 0/1; priority enumeration not established.'))
define([27], 'Alde 3020+', 4,
       (0xC0, 'Electric selection', 'Raw selection code 0..3.'), unknown(0x3F))
define([27], 'Alde 3020+', 5,
       flag(0x01, 'Heater state'), flag(0x02, 'Manual mode'), flag(0x04, 'Error'),
       (0x18, 'Water selection', 'Raw 0..3; corresponding control uses 0=off, 1=normal, 2=boost.'),
       flag(0x80, 'Pump'), unknown(0x60))

# ATC bit fields are known; the enum labels are not.
define([36], 'AL-KO ATC', 0,
       (0x07, 'ATCStatus', 'Raw 0..7; enumeration not supplied.'),
       (0x18, 'ATCHistogram', 'Raw 0..3; enumeration not supplied.'),
       (0xE0, 'ATCACCStatus', 'Raw 0..7; enumeration not supplied.'))

# Water and space heating.
define([56], 'Truma CP+', 5,
       (0x80, 'Heater error', 'Set flag triggers Truma diagnostic request on ID 60.'), unknown(0x7F))
define([57, 58], 'Truma CP+', 1,
       (0x0F, 'Target high bits', 'High four bits of 12-bit temperature; see combined target row.'),
       (0xF0, 'TrumaMode', 'Raw 0..15; enumeration not established. Retained from ID 58 by the control path.'))
define([55, 57], 'Whale', 2,
       (0xFF, 'Energy selection (overlaps subfields)', '0=off; 1=electric stage 1; 3=stage 2; 7=stage 3; 16=gas; 17/19/23=gas + electric stages 1/2/3.'),
       flag(0x10, 'Gas'),
       (0x07, 'Electric stage code', '0=none; 1=stage 1; 3=stage 2; 7=stage 3. Encoded selection, not independent switches.'),
       unknown(0xE8, 'Not set in the listed energy codes; wider meaning not established.'))
define([57], 'Whale', 3,
       (0xFF, 'Feature code (overlaps flags)', '0=normal; 4=frost (5 C); 2=night (16 C); 1=fan only (0 C).'),
       flag(0x04, 'Frost'), flag(0x02, 'Night'), flag(0x01, 'Fan only'),
       unknown(0xF8, 'Not set in listed feature codes; wider meaning not established.'))
for ident, prefix in [(56, 'HotWater'), (58, 'Heating')]:
    define([ident], 'Whale', 4,
           (0xFF, 'Supply voltage', 'Inferred voltage=B4/10 V: source compares >90 with >9 V. Not independently verified.', 'Inferred from source comparison'))
    define([ident], 'Whale', 5,
           (0x0F, prefix + 'Error', 'Raw 0..15; error enumeration not established. Updated 2026 in source.'), unknown(0xF0))
    define([ident], 'Whale', 7,
           (0xFF, prefix + 'Status (overlaps flag)', 'Raw status byte; only bit 6 is identified here.'),
           flag(0x40, 'Panel busy / manual interaction'), unknown(0xBF))
define([58], 'Eberspacher', 6,
       (0xF0, 'Heater mode', '0=off; 1=on TT; 2=ventilate; 4=purge; 5=on TTCT (source labels).'), unknown(0x0F))
define([58], 'Eberspacher', 7, flag(0x80, 'EbLinStatus'), unknown(0x7F))

# Webasto: source trace in docs/EC600-WEBASTO-LIN.md; every bit is accounted for.
define([57], 'Webasto', 0,
       (0x07, 'Attempted requested mode; actual zero', 'M & 7 is ANDed with target bits in b7..5; result zero. Source requested modes: 0=off, 1=park heating, 2=not used, 3=ventilation, 4=boost, 5=eco.', 'Source conflict'),
       (0x18, 'Fixed zero', 'Both bits written zero; protocol meaning unknown.'),
       (0xE0, 'Attempted target low bits; actual zero', '(T & 7) << 5 is ANDed with mode in b2..0; result zero.', 'Source conflict'))
define([57], 'Webasto', 1,
       (0x07, 'Surviving target bits 5..3', '((T & 248) >> 3) AND (A & 7); A=255, so (T >> 3) & 7. Intended T=5..35 C yields 0..4. Not a complete temperature.', 'Source conflict'),
       (0xF8, 'Fixed zero', 'Upper bits cleared; intended altitude placement unknown.', 'Source conflict'))
define([57], 'Webasto', 2,
       (0x1F, 'Altitude upper part', '(A & 248) >> 3 = 31 because A=255 (signal not available comment).'),
       (0x20, 'Fixed zero', 'Bit 5 is zero.'),
       (0xC0, 'Fixed ones', 'Adds 0xC0. Whole byte is 0xDF.'))
define([57], 'Webasto', 4,
       (0x0F, 'Power low nibble', 'Combine with B5 bits 2..0. Source comment: requested power 0..100 means 0..100%; no clamp here.'),
       (0xF0, 'Fixed upper nibble', 'Extracted value=15 (0xF).'))
define([57], 'Webasto', 5,
       (0x07, 'Power bits 6..4', 'Combine with B4 bits 3..0; input power bit 7 discarded.'),
       (0xF8, 'Fixed zero', 'Upper five bits are written zero.'))
define([58], 'Webasto', 0,
       (0x1F, 'Working mode', 'WebastoCMode=B0 & 31. Stored only; raw 0..31; enum unknown and differs from requested modes.'),
       (0x80, 'Diagnostic flag', 'WebastoCDiag=(B0 & 128) >> 7. Stored only; no diagnostic action follows.'), unknown(0x60))
define([58], 'Webasto', 1,
       (0x03, 'Error status', 'WebastoCErrorStatus=B1 & 3. Stored only; raw 0..3; enum unknown.'),
       (0xFC, 'Temperature fragment 1', 'WebastoCTemperature=(B1 & 252) AND (B2 & 3)=0; stored only. Parser says actual cabin temperature, declaration says medium temperature. Scale unknown.', 'Source conflict'))
define([58], 'Webasto', 2,
       (0x03, 'Temperature fragment 2', 'Final WebastoCTemperature=0; stored only. Raw extraction here is not the final source result.', 'Source conflict'),
       (0xFC, 'Voltage fragment 1', 'WebastoCVoltage=(B2 & 252) AND (B3 & 3)=0; stored only. Scale unknown.', 'Source conflict'))
define([58], 'Webasto', 3,
       (0x03, 'Voltage fragment 2', 'Final WebastoCVoltage=0; stored only. Raw extraction here is not the final source result.', 'Source conflict'), unknown(0xFC))
define([58], 'Webasto', 6,
       (0x08, 'Bit masked as heater error', 'WebastoCHeaterError=(B6 & 8) >> 4=0; stored only. Raw extraction here is not the final source result.', 'Source conflict'),
       (0x10, 'Bit masked as altitude', 'WebastoCAltitude=(B6 & 16) >> 5=0; stored only. Raw extraction here is not the final source result.', 'Source conflict'), unknown(0xE7))
define([58], 'Webasto', 7,
       (0x40, 'Comp Error (source label)', 'WebastoCCompError=(B7 & 64) >> 6; stored only. Expansion of Comp is unknown.'),
       (0x80, 'LIN error', 'WebastoCLinError=(B7 & 128) >> 7. Accepted reply: 1 clears LinStatusHeat, 0 sets it. Invalid receive also clears it. Reported at CAN 53 B6 b0.'), unknown(0x3F))

# Compound values supplement the individual byte fields.
combined = {}
def combine(ids, device, byte, bits, mask, name, expression, values, evidence='Source-derived'):
    for ident in ids:
        combined.setdefault((ident, device), []).append([byte, bits, mask, name, expression, values, evidence])

for ident, low in [(8, 0), (23, 2)]:
    combine([ident], 'Truma aircon', f'B{low}/B{low+1}', '16-bit little-endian', '0xFFFF', 'Combined target temperature',
            f'B{low} | (B{low+1} << 8)', 'degrees C=(raw-2730)/10.')
for ident in (8, 23):
    combine([ident], 'Dometic aircon', 'B0/B1', 'B0 b7,b1..0; B1 b1..0', 'B0:0x83; B1:0x03', 'Combined mode',
            '((B0 & 128) >> 3) | (B1 & 3) | ((B0 & 3) << 2)',
            'Composite code; enumeration not established. Assembly documented by the information path; do not interpret one mode bit alone.')
combine([26, 27], 'Alde / EC645 legacy', 'B1/B2', '16-bit little-endian', '0xFFFF', 'Combined room target',
        'B1 | (B2 << 8)', 'degrees C=raw/10. ID 27 target bytes are compared with ID 26.')
combine([55, 56], 'Truma CP+', 'B0/B1', '16-bit little-endian', '0xFFFF', 'Combined water target',
        'B0 | (B1 << 8)', '0x0000=off; 0x0C3A=40 C Eco; 0x0D02=60 C Hot. Byte order on wire: 00 00; 3A 0C; 02 0D.')
combine([55, 56], 'Truma CP+', 'B3/B4', '16-bit little-endian', '0xFFFF', 'Combined electric power',
        'B3 | (B4 << 8)', '0x0384=900 W; 0x0708=1800 W. Other supported values not established.')
combine([55, 57], 'Whale', 'B0/B1', '16-bit little-endian', '0xFFFF', 'Combined target temperature',
        'B0 | (B1 << 8)', 'degrees C=(raw-2730)/10. ID 55 modes: off=0 C (0x0AAA), frost=25 C, Eco=55 C, Max=72 C. ID 57 feature modes are in B3.')
combine([57], 'Whale', 'B6/B7', '16-bit little-endian', '0xFFFF', 'Combined room temperature',
        'B6 | (B7 << 8)', 'degrees C=(raw-2730)/10; forwarded from IntTL/IntTH.')
combine([57, 58], 'Truma CP+', 'B0/B1', 'B0 b7..0; B1 b3..0', '0x0FFF', 'Combined room target',
        'B0 | ((B1 & 0x0F) << 8)', 'raw=0 means off; otherwise degrees C=(raw-2730)/10. Exclude TrumaMode upper nibble.')
combine([57], 'Eberspacher', 'B3/B4', '16-bit little-endian', '0xFFFF', 'Combined room temperature raw',
        'B3 | (B4 << 8)', 'Source encodes raw=IntTempByte*10+500. IntTempByte scale is not established; no physical conversion assigned.')
combine([57], 'Webasto', 'B4/B5', 'B4 b3..0; B5 b2..0', '0x007F', 'Combined power',
        '(B4 & 15) | ((B5 & 7) << 4)', 'Represented raw 0..127; source labels intended 0..100 as percent. Input comes from SetOut3[1] when on, otherwise zero; no clamp in Webasto routine.')
combine([61], 'Truma diagnostic', 'B4/B5', 'Class and code bytes', 'Context-dependent', 'PSU packed error',
        '255 if B5 == 255 else B5 + (B4 << 5)', 'Applies only to the matching Truma diagnostic response. Error enumeration not established.')


def pid(ident):
    b = lambda n: (ident >> n) & 1
    return ident | ((b(0)^b(1)^b(2)^b(4)) << 6) | ((1^b(1)^b(3)^b(4)^b(5)) << 7)


def evidence(text):
    t = text.lower()
    if any(s in t for s in ('suspect', 'disjoint', 'outside byte range', 'yields zero', 'produces zero')):
        return 'Source conflict'
    if 'disabled' in t:
        return 'Source-derived; disabled branch noted'
    if any(s in t for s in ('not defined', 'not decoded', 'no meaning', 'no response-field meaning', 'no local meaning')):
        return 'Unknown'
    if any(s in t for s in ('not established', 'uncertain', 'not assumed', 'not supplied')):
        return 'Source-derived; incomplete meaning'
    if any(s in t for s in ('request-specific', 'transaction', 'depending on frame', 'not assigned universally')):
        return 'Context-dependent'
    if 'comment' in t or 'rfu' in t:
        return 'Source comment; not independently verified'
    return 'Source-derived'


def bit_label(mask):
    groups = []
    bits = [i for i in range(7, -1, -1) if mask & (1 << i)]
    while bits:
        high = low = bits.pop(0)
        while bits and bits[0] == low - 1:
            low = bits.pop(0)
        groups.append(f'b{high}' if high == low else f'b{high}..{low}')
    return ', '.join(groups)


def extract(byte, mask):
    shift = (mask & -mask).bit_length() - 1
    # Non-contiguous masks stay in place; compact fields shift to bit zero.
    contiguous = ((mask >> shift) & ((mask >> shift) + 1)) == 0
    return f'B{byte}' if mask == 255 else (f'(B{byte} & 0x{mask:02X}) >> {shift}' if shift and contiguous else f'B{byte} & 0x{mask:02X}')


# Human-readable names for whole-byte fields; packed fields have their own names.
byte_names = {}
def names(ids, device, labels):
    labels = labels.split('|')
    assert len(labels) == 8
    for ident in ids:
        byte_names[ident, device] = labels

names([8], 'Truma aircon', 'Target low|Target high|Fan code|Operating mode|AC power low|AC power high|Light level|RFU (source comment)')
names([23], 'Truma aircon', 'Unknown|Unknown|Target low|Target high|Fan code|Operating mode|AC power low|AC power high')
names([8, 23], 'Dometic aircon', 'On / light / mode|Target / fan / mode|Unknown / synchronised data|Unknown / synchronised data|iFeel temperature|Light dimming data|Unknown / synchronised data|Synchronisation / change flags')
names([11], 'Dometic fridge', 'Operating mode|Cooling setting|Frame-heater setting|Control constant|Control constant|Availability flags|Control spare|Control spare')
names([12], 'Dometic fridge', 'Operating mode|Cooling level|Frame-heater setting|Fridge error code|Fridge status flags|Stored / echoed data|Stored / echoed data|Stored / echoed data')
names([26], 'Alde 3020+', 'Zero constant|Zero constant|Zero constant|Target / gas|Electric / HeatingNow|On / water selection|Zero constant|Zero constant')
names([27], 'Alde 3020+', 'Room temperature|Bed temperature (comment)|Outdoor temperature (comment)|Target / gas / priority|Electric selection|Heater / water / pump flags|Stored data|Stored data')
names([26], 'Alde / EC645 legacy', 'Settings|Target low|Target high|Normal control constant|Zero constant|Zero constant|Zero constant|Zero constant')
names([27], 'Alde / EC645 legacy', 'Settings compared with control|Target low|Target high|HeatingStatus|HeatingError|Unknown|Unknown|EC645 CRC checking enabled')
names([26], 'Alde / EC645 clock', 'Zero constant|Zero constant|Zero constant|Clock marker / day|Hour|Minute|Zero constant|Zero constant')
names([36], 'AL-KO ATC', 'ATC status fields|Unknown|Unknown|Unknown|Unknown|Unknown|Unknown|Unknown')
names([36], 'LEVC experiment (disabled)', 'LEVC_Grid (disabled)|LEVC_DCDC (disabled)|Unknown|Unknown|Unknown|Unknown|Unknown|LEVC_Brake (disabled)')
names([55], 'Truma CP+', 'Water target low|Water target high|Energy code|Electric power low|Electric power high|Zero constant|Zero constant|Zero constant')
names([56], 'Truma CP+', 'Water target low|Water target high|Energy code|Electric power low|Electric power high|Heater error flag|Unknown|Unknown')
names([55], 'Whale', 'Water target low|Water target high|Energy code|Forwarded data|Forwarded data|Forwarded data|Forwarded data|Forwarded data')
names([56], 'Whale', 'Unknown|Unknown|Unknown|Unknown|Supply voltage|HotWaterError|Unknown|HotWaterStatus')
names([57], 'Truma CP+', 'Room target low|Target high / mode|Zero constant|Energy code|Zero constant|Zero constant|Zero constant|Zero constant')
names([58], 'Truma CP+', 'Room target low|Target high / mode|Unknown|Unknown|Energy setting|Unknown|Unknown|Unknown')
names([57], 'Whale', 'Target low|Target high|Energy code|Feature code|Forwarded data|Forwarded data|Room temperature low|Room temperature high')
names([58], 'Whale', 'Unknown|Unknown|Unknown|Unknown|Supply voltage|HeatingError|Unknown|HeatingStatus')
names([57], 'Eberspacher', 'Fixed override|Command mode|Target temperature|Room temperature raw low|Room temperature raw high|Zero when power on|Zero when power on|Zero when power on')
names([58], 'Eberspacher', 'Unknown|HeatingError|Unknown|Unknown|Target temperature|Unknown|Heater mode|EbLinStatus')
names([57], 'Webasto', 'Suspect mode / target packing|Suspect target / altitude packing|Altitude upper part|Fixed 0xFF|Power low / fixed nibble|Power high|Fixed 0xFF|Advised cabin temperature')
names([58], 'Webasto', 'Working mode / diagnostic|Error / temperature fragment|Temperature / voltage fragments|Voltage fragment|Unknown|Unknown|Suspect heater error / altitude|Comp / LIN errors')
names([60], 'Diagnostic / configuration', 'Address / NAD|PCI / length control|Service / length|Request-specific data|Request-specific data|Request-specific data|Request-specific data|Request-specific data / new NAD')
names([61], 'Truma diagnostic', 'Addressing / transaction|Transport data|Service / data|Response data|Truma error class|Truma error code|Unknown|Unknown')
names([61], 'Alde diagnostic / RCP', 'Device NAD|Transport / sequence|Length / RCP data|Service / RCP data|Transaction-specific data|RCP payload start|RCP payload|RCP payload')
names([61], 'FreshJet diagnostic', 'Address comparison|Length comparison|Suspect response comparison|Unknown|Unknown|Unknown|Unknown|Unknown')

records = []
for entry in catalogue['entries']:
    for layout in entry['layouts']:
        name = layout['name']
        context = ('Disabled experiment' if 'disabled' in name else
                   'Clock setting' if name.endswith('clock') else
                   'Diagnostic request' if entry['id'] == 60 else
                   'Diagnostic response' if entry['id'] == 61 else
                   'Control' if entry['direction'].startswith('PSU') else 'Status')
        applicability = ('Disabled / historical' if 'disabled' in name else
                         'Source conflict / manual selection' if name == 'Webasto' else
                         'Transaction-dependent' if entry['id'] >= 60 else
                         'Legacy interface only' if 'EC645' in name else 'Selected device/layout only')
        notes = ' '.join(s for s in (entry['notes'], layout['notes']) if s)
        records.append((entry, layout, context, applicability, notes, entry['id']))
        if entry['id'] == 11:
            echo = dict(entries[12]['layouts'][0])
            echo['source'] += '; FCM_FridgeCTRL (49420) for echo'
            records.append((entry, echo, 'Status echo', applicability,
                            'PSU echoes ID 12 when no settings change is pending. Use status meanings for all eight bytes. ' + notes, 12))

frame_rows, byte_rows, signal_rows = [], [], []
for entry, layout, context, applicability, notes, definition_id in records:
    ident, device, source = entry['id'], layout['name'], layout['source']
    frame_rows.append([ident, f'0x{ident:02X}', f'0x{pid(ident):02X}', device, context, entry['direction'],
                       entry['length'], entry['checksum'], 'PSU', applicability, entry['name'],
                       'Not established', notes or 'No additional note.', source])
    for byte, text in enumerate(layout['bytes']):
        if ident == 11 and context == 'Control':
            text = {3: 'Written as 0. Old hour-setting code disabled.',
                    4: 'Written as 0. Old minute-setting code disabled.',
                    6: 'Written as 0 (spare) in control frame.',
                    7: 'Written as 0 (spare) in control frame.'}.get(byte, text)
        ev = evidence(text)
        if applicability == 'Disabled / historical':
            ev = 'Disabled / historical; ' + ev
        byte_rows.append([ident, device, context, f'B{byte}', entry['direction'], text, ev, notes, source])
        definitions = fields.get((definition_id, device, byte))
        if definitions is None:
            definitions = [(255, byte_names[definition_id, device][byte], text, ev)]
        coverage = 0
        for item in definitions:
            mask, signal, values = item[:3]
            coverage |= mask
            field_ev = item[3] if len(item) > 3 else evidence(values)
            if applicability == 'Disabled / historical':
                field_ev = 'Disabled / historical; ' + field_ev
            signal_rows.append([ident, device, context, f'B{byte}', bit_label(mask), f'0x{mask:02X}', signal,
                                extract(byte, mask), values, field_ev, source, text, notes])
        assert coverage == 255, (ident, device, byte, 'Unaccounted bits')
    for row in combined.get((definition_id, device), []):
        signal_rows.append([ident, device, context] + row + [source, 'Combined field; see constituent bytes in Byte Map.', notes])

layout_count = sum(len(e['layouts']) for e in catalogue['entries'])
guide = [
    ['Coverage', f'All {len(entries)} LIN IDs and {layout_count} ID/device layouts in the reviewed catalogue, plus the separate ID 11 refrigerator status-echo context. Includes diagnostic, legacy and disabled experiment entries with explicit labels.'],
    ['Start here', 'Byte Map is the familiar one-row-per-byte overview. Filter Device / Variant, LIN ID and Context together. Signals gives bit masks, value meanings and combined multi-byte conversions. Frames retains applicability and layout notes.'],
    ['Direction', 'The PSU sends every LIN header. Data direction describes who supplies the payload. Appliance → PSU therefore still begins with a PSU header. All listed transactions use 8 data bytes, with checksum separate.'],
    ['Numbering', catalogue['byteNumbering'] + ' Hex values use a 0x prefix. Expected PID is the protected identifier, not the six-bit LIN ID. Bit ranges include both endpoints.'],
    ['Expressions', 'Extraction uses bitwise & (AND), | (OR), << (left shift) and >> (right shift). Masks on separate bytes are stated explicitly. Expressions are documentation text, not executable Excel formulas. Non-contiguous masks remain in their original positions.'],
    ['Evidence', 'Source-derived describes the reviewed PSU implementation, not a manufacturer-certified protocol or a verified physical capture. Unknown fields, missing enums, inferred voltage scaling and source conflicts stay explicit. Amber cells highlight limitations.'],
    ['Shared heater IDs', 'IDs 55–58 have different layouts by make. Select the installed variant; a valid checksum or an ID alone does not identify the make. Truma status energy is B4 on ID 58, while control energy is B3 on ID 57.'],
    ['Device selection', 'Existing analyzer auto selection uses confirmed-format CAN 133 B2: 2=Truma CP+, 3=Whale, 4=Eberspacher. This describes configuration, not physical presence. Alde generation is manual. Setting 5 is called Webasto by firmware and Whale Ci-Bus by the CAN table; it remains ambiguous/manual.'],
    ['Fridge echo', 'ID 11 sends control data when a change is pending, otherwise echoes ID 12 status. The Control and Status echo contexts are separate, especially for B3–B5. Thetford remains a placeholder without a verified layout.'],
    ['Alde variants', '3020+, EC645 legacy control/status, and EC645 clock have separate layouts. Clock frames must not be decoded as ordinary control. Legacy B7 CRC-enable is payload data, not the separate LIN checksum.'],
    ['Diagnostics', 'IDs 60/61 use classic checksum. Field meanings depend on the addressed device and transaction. Alde RCP requires multi-frame reconstruction; no universal single-frame bit dictionary is implied. NAD alone may be ambiguous.'],
    ['Source conflicts', 'Truma aircon status uses executed offsets rather than contradictory comments. Webasto contains suspect AND/shift operations; no corrected manufacturer protocol is invented. FreshJet checks an 8-bit byte against 0x1F2, which cannot be an expected byte.'],
    ['Disabled example', 'LEVC is a disabled historical experiment, not the active AL-KO ATC layout. Its rows remain available for completeness and are labelled Disabled experiment.'],
    ['Incomplete reception', 'A header without a reply is not an eight-byte zero payload. Keep raw data, PID/checksum validity and receive status when applying this reference to a capture.'],
    ['Timing', 'Frame periods are not established in the reviewed catalogue. The active main loop schedules selected devices; no fixed millisecond intervals are assumed.'],
    ['Examples', 'Examples are synthetic calculations or explicitly identified source examples. They are not observed traffic, approved transmit templates, or proof that a device accepts a complete payload. Unknown bytes are never filled in to suggest a verified message.'],
    ['Source', catalogue['source'] + '; catalogue: lin/ec600-lin-scheme.json; reference: docs/EC600-LIN-REFERENCE.md; findings: docs/EC600-LIN-FINDINGS.md.'],
    ['Source SHA-256', catalogue['sha256']],
    ['Applicability', 'Based on the existing review of matching V57B generated C and Flowcode. Installed firmware is not identified by this workbook. Historical firmware variants were not exhaustively catalogued.'],
    ['Regenerate', 'Requires Python and openpyxl 3.1+. Run python3 scripts/export-lin-full-workbook.py. The exporter verifies the catalogue hash so curated bit definitions cannot silently drift after catalogue edits.'],
]
sheet('Read Me', 'EC600 LIN | Full information table',
      f'{len(entries)} IDs • {layout_count} catalogue layouts + fridge echo • Filter by Device / Variant and Context',
      ['Topic', 'Explanation'], guide, [29, 125])
sheet('Frames', 'FRAME SUMMARY | All reviewed LIN layouts',
      'Every row is one ID / device / context. Shared IDs are alternative layouts. Expected PID includes parity; the checksum is additional to payload length.',
      ['LIN ID dec', 'LIN ID hex', 'Expected PID', 'Device / Variant', 'Context', 'Data direction', 'Payload bytes', 'Checksum', 'Header sender', 'Applicability', 'Frame description', 'Period (ms)', 'Layout notes / limitations', 'Source routines / lines'],
      frame_rows, [12, 13, 14, 27, 22, 21, 12, 14, 14, 30, 44, 20, 80, 48])
sheet('Byte Map', 'BYTE MAP | All devices and variants',
      'One row per payload byte. Filter Device / Variant and Context before interpreting a shared ID. Full source descriptions and layout qualifications are retained.',
      ['LIN ID dec', 'Device / Variant', 'Context', 'Byte', 'Data direction', 'Function / values', 'Evidence', 'Layout notes / limitations', 'Source routines / lines'],
      byte_rows, [12, 27, 22, 9, 21, 85, 30, 72, 48])
sheet('Signals', 'SIGNAL DICTIONARY | Bits, values and conversions',
      'One row per field, with extra rows for combined values. Overlapping whole-byte/bit fields are intentional. Unknown bits are retained. Source conflicts are labelled.',
      ['LIN ID dec', 'Device / Variant', 'Context', 'Byte(s)', 'Bits / ordering', 'Mask hex', 'Signal / variable', 'Extract raw value', 'Options / conversion / notes', 'Evidence', 'Source routines / lines', 'Source byte description', 'Layout notes / limitations'],
      signal_rows, [12, 27, 22, 12, 30, 25, 33, 46, 82, 32, 48, 72, 72])

# Partial examples avoid inventing values for undocumented payload bytes.
examples = []
def example(ident, device, context, data, explanation, kind='Synthetic field example'):
    examples.append([ident, device, context, kind, data, explanation])
example(8, 'Truma aircon', 'Control', 'B0/B1=72 0B; B2=72; B3=05',
        '0x0B72=2930 → (2930-2730)/10=20 C. Fan 0x72=medium. B3 low nibble 5=cooling. Other bytes unspecified.')
example(23, 'Truma aircon', 'Status', 'B2/B3=72 0B; B4=72; B5=05',
        'Same target/fan/mode example, using executed status offsets B2/B3, B4 and B5. Do not move target to B0/B1.')
example(8, 'Dometic aircon', 'Control', 'B1=40; B7=04',
        'B1 upper nibble=4 → target 20 C; fan field=0 (enum unknown). B7 bit 2=1 is sync. Composite operating mode also needs B0; no mode inferred.')
example(11, 'Dometic fridge', 'Control', '03 03 00 00 00 10 00 00',
        'Gas; on=1; cooling raw 3 → PSU setting 2. B5=0x10: auto available=1, frame heater fitted=0. Remaining zero values follow this example control layout.', 'Synthetic full payload')
example(12, 'Dometic fridge', 'Status', '03 03 00 00 18 00 00 00',
        'Gas; cooling raw 3 → PSU setting 2. B4=0x18: door=1, manual=1, FridgeState=0 (enum unknown). B3 error code 0 and B5–B7 meanings are not assigned.', 'Synthetic full payload')
example(11, 'Dometic fridge', 'Status echo', '03 03 00 00 18 00 00 00',
        'PSU echo of the preceding status example. B4 remains status flags. Do not apply control meanings to B5.', 'Synthetic full payload')
example(26, 'Alde 3020+', 'Control', 'B3=5E; B5=09',
        'B3 low six bits=30 → target 20 C, gas flag=1. B5 on=1 and water selection=1 (normal). B4 electric selection not specified.')
example(27, 'Alde 3020+', 'Status', 'B0=7C; B3=5E',
        'Room temperature: (124-84)/2=20 C. Target field=30 → (30+10)/2=20 C; gas flag=1.')
example(26, 'Alde / EC645 legacy', 'Control', 'B0=2B; B1/B2=C8 00; B3=40',
        'Heater=1; gas=1; electric field=2 kW; water field=2 (normal). Target 0x00C8=200 → 20 C. B3=0x40 ordinary control marker.')
example(26, 'Alde / EC645 clock', 'Clock setting', 'B3=80; B4=0E; B5=1E',
        'Clock marker=1; encoded day field=0 → source Day=1 (weekday convention unknown). Hour=14; minute=30.')
example(36, 'AL-KO ATC', 'Status', 'B0=AB',
        'Binary 10101011: ATCStatus=3, ATCHistogram=1, ATCACCStatus=5. Enumeration meanings remain unknown.')
example(55, 'Truma CP+', 'Control', 'B0/B1=3A 0C; B2=02; B3/B4=84 03',
        'Water Eco 40 C; energy=2 electric; power 0x0384=900 W. B0/B1 all-zero would instead mean off.')
example(56, 'Truma CP+', 'Status', 'B5=80',
        'Error bit 7 set; source triggers a diagnostic request on ID 60. Other bits and bytes are not inferred.')
example(55, 'Whale', 'Control', 'B0/B1=D0 0C; B2=13',
        '0x0CD0=3280 → 55 C Eco. 0x13=19 → gas + electric stage 2. Off temperature is AA 0A (0 C), not 00 00.')
example(57, 'Truma CP+', 'Control', 'B0/B1=72 2B; B3=03',
        'Target uses B1 low nibble only: raw 0x0B72=2930 → 20 C. TrumaMode=2 (enum unknown). Energy=3 mixed.')
example(58, 'Truma CP+', 'Status', 'B0/B1=72 2B; B4=03',
        'Target 20 C; TrumaMode=2 (enum unknown). Status energy is B4=3 mixed, not control byte B3.')
example(57, 'Whale', 'Control', 'B0/B1=4A 0B; B2=10; B3=02; B6/B7=72 0B',
        'Target raw 2890 → 16 C; energy 16=gas; feature 2=night; room raw 2930 → 20 C. B4/B5 unspecified.')
example(56, 'Whale', 'Status', 'B4=78; B5=03; B7=40',
        'Supply raw 120 implies 12.0 V from source comparison. HotWaterError=3 (enum unknown); panel busy bit=1. ID 58 uses HeatingError with the same masks.')
example(57, 'Eberspacher', 'Control', 'B0=FF; B1=08; B2=50',
        'Override=0xFF; command 8=ventilate; target raw 80 → (80-40)/2=20 C. Room raw conversion is not inferred.')
example(58, 'Eberspacher', 'Status', 'B4=50; B6=40; B7=80',
        'Target 20 C; status mode 4=purge; EbLinStatus=1. Status mode codes differ from command codes.')
example(57, 'Webasto', 'Control', 'B4=F5; B5=02; B7=46',
        'Combined power=(5 | (2<<4))=37; source labels this requested percent. Cabin temperature=70-50=20 C. This does not repair suspect B0/B1 packing.')
example(58, 'Webasto', 'Status', 'B6=18',
        'Raw bits 3 and 4 are set. The reviewed source shifts these masked values by 4 and 5, respectively, producing zero. No corrected device semantics asserted.')
example(60, 'Diagnostic / configuration', 'Diagnostic request', '02 06 B2 23 17 46 40 13',
        'Truma diagnostic source example: NAD 0x02; PCI 0x06; service 0xB2. Request context is required. Classic checksum is separate.', 'Source example; not a capture')
example(60, 'Diagnostic / configuration', 'Diagnostic request', '10 02 A0 01 FF FF FF FF',
        'Alde RCP source example: NAD 0x10; PCI 0x02; service 0xA0; request data 0x01. Response may require multi-frame reconstruction.', 'Source example; not a capture')
example(61, 'Truma diagnostic', 'Diagnostic response', 'B4=02; B5=03',
        'PSU packed error=3+(2<<5)=67. If B5=255, PSU keeps 255 directly. Error text/enum unknown; applies only to Truma response context.')
example(61, 'Alde diagnostic / RCP', 'Diagnostic response', 'First frame: B5..B7; following frames: B2..B7',
        'PSU appends these fragments to RCPString, then interprets tagged elements. No invented complete payload or single-frame signal dictionary.', 'Source layout explanation')
example(61, 'FreshJet diagnostic', 'Diagnostic response', 'B2 compared with 0x1F2 in source',
        '0x1F2 exceeds one byte. It is a suspect source comparison, not a legal expected B2 value.', 'Source conflict explanation')
example(36, 'LEVC experiment (disabled)', 'Disabled experiment', 'B0 / B1 / B7',
        'Named LEVC_Grid / LEVC_DCDC / LEVC_Brake only in the disabled experiment. No numeric values or active ATC meanings inferred.', 'Disabled source explanation')
sheet('Examples', 'WORKED EXAMPLES | All device families',
      'Hex byte values. Partial examples leave other bytes unspecified. Full payloads are synthetic or explicitly labelled source examples. No captured traffic is claimed.',
      ['LIN ID dec', 'Device / Variant', 'Context', 'Example evidence', 'Example bytes / fields', 'Decoded explanation'],
      examples, [12, 27, 22, 31, 59, 113])

# Keep the familiar view first on opening; headings and first three columns freeze.
wb.active = 2
for ws in wb:
    ws.sheet_properties.tabColor = NAVY if ws.title != 'Read Me' else '2A8D83'
wb.save(OUTPUT)

# Validate the delivered workbook, not just the input lists.
check = load_workbook(OUTPUT)
assert check.sheetnames == ['Read Me', 'Frames', 'Byte Map', 'Signals', 'Examples']
assert {r[0] for r in frame_rows} == set(entries)
assert len(frame_rows) == layout_count + 1
assert len(byte_rows) == len(records) * 8
for entry, layout, context, _, _, _ in records:
    key = (entry['id'], layout['name'], context)
    actual = [r[3] for r in check['Byte Map'].iter_rows(min_row=5, values_only=True) if r[:3] == key]
    assert actual == [f'B{i}' for i in range(8)], key
    assert any(r[:3] == key for r in check['Signals'].iter_rows(min_row=5, values_only=True)), key
assert all(key[0] in entries for key in fields)
for ws in check:
    assert len(ws.tables) == 1 and ws.freeze_panes
    assert all(cell.data_type != 'f' for row in ws for cell in row), 'Descriptions must not become Excel formulas'
for r in examples:
    if r[3] in ('Synthetic full payload', 'Source example; not a capture'):
        assert len(bytes.fromhex(r[4])) == 8
assert pid(11) == 0x8B and pid(12) == 0x4C and pid(58) == 0xBA and pid(61) == 0x7D
print(f'Created and reopened {OUTPUT}')
print(f'{len(entries)} IDs; {layout_count} catalogue layouts + status echo; {len(frame_rows)} frame rows; {len(byte_rows)} byte rows; {len(signal_rows)} signal rows; {len(examples)} examples.')
