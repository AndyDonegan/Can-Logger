#!/usr/bin/env python3
"""Build the proposed Dometic fridge LIN workbook. Requires openpyxl (3.1+)."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs/LIN-Table-Layout-Dometic-Fridge.xlsx'
catalogue = json.loads((ROOT / 'lin/ec600-lin-scheme.json').read_text())
entries = {e['id']: e for e in catalogue['entries']}
wb = Workbook()
wb.remove(wb.active)
NAVY, TEAL, AMBER, GREY = '17324D', 'DDF2EF', 'FFF0CC', 'EDF0F4'


def sheet(name, title, subtitle, headers, rows, widths):
    ws = wb.create_sheet(name)
    end = get_column_letter(len(headers))
    for n, text in [(1, title), (2, subtitle)]:
        ws.merge_cells(f'A{n}:{end}{n}')
        cell = ws.cell(n, 1, text)
        cell.fill = PatternFill('solid', fgColor=NAVY if n == 1 else TEAL)
        cell.font = Font(name='Calibri', size=18 if n == 1 else 11,
                         bold=n == 1, color='FFFFFF' if n == 1 else NAVY)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 44
    for col, text in enumerate(headers, 1):
        c = ws.cell(4, col, text)
        c.font = Font(name='Calibri', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 32
    for r, row in enumerate(rows, 5):
        for col, value in enumerate(row, 1):
            c = ws.cell(r, col, value)
            c.font = Font(name='Calibri', size=11, color=NAVY)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(value, str) and ('Unknown' in value or 'not established' in value):
                c.fill = PatternFill('solid', fgColor=AMBER)
        lines = max(max(len(str(v)) // max(int(widths[i]) - 3, 1) + 1,
                        str(v).count('\n') + 1) for i, v in enumerate(row))
        ws.row_dimensions[r].height = max(36, min(150, 16 * (lines + 1)))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    table = Table(displayName=name.replace(' ', '') + 'Table', ref=f'A4:{end}{ws.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'D5' if len(headers) > 5 else 'A5'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{end}{ws.max_row}'
    ws.oddFooter.center.text = 'Dometic fridge LIN | Proposed layout | Page &P of &N'
    return ws


guide = [
    ['Start here', 'Proposed Excel layout for one device: Dometic refrigerator with the EC600 PSU. Start with Byte Map for the familiar CAN-style overview; use Signals for bit-level details.'],
    ['Workbook tabs', 'Frames: transaction summary. Byte Map: one row per payload byte. Signals: one row per field, with mask and values. Examples: worked synthetic payloads. All data tables have filters and frozen headings.'],
    ['Direction', 'Direction describes the payload publisher and receiver. The PSU sends every LIN header, including when the fridge supplies the response. TX/RX alone would be ambiguous for a passive analyzer.'],
    ['Numbering', 'B0 is the first payload byte; B7 the eighth. b0 is the least significant bit. Hex values have a 0x prefix. Bits b2..0 means bits 2, 1 and 0. PID includes parity and is separate from the six-bit LIN ID.'],
    ['Critical layout distinction', 'ID 11 (0x0B) is a control frame when settings change. Otherwise the PSU echoes the ID 12 status payload. Select the Control or Status echo context before interpreting bytes, especially B3, B4 and B5.'],
    ['Evidence', 'Source-derived means documented in the existing reviewed PSU catalogue; it does not mean verified on a physical fridge. Unknown means the reviewed source does not establish a meaning. Amber cells highlight gaps.'],
    ['Unknown is not unused', 'Unknown bytes/bits remain present in the payload. A value written as zero in a control branch does not establish that the corresponding status field is unused.'],
    ['Examples', 'All example payloads are constructed illustrations, not captured traffic or transmit templates. No timing, checksum validity or physical device behaviour has been verified by these examples.'],
    ['Scope', 'Dometic fridge layout only. Thetford is a placeholder in the source and is not covered. No device model number or installed firmware revision is inferred.'],
    ['Source', catalogue['source'] + '; reviewed mapping: lin/ec600-lin-scheme.json; docs/EC600-LIN-REFERENCE.md, IDs 11 and 12.'],
    ['Source SHA-256', catalogue['sha256']],
    ['Source routines', 'Control: FCM_Fridge6 (32313), FCM_FridgeCTRL (49420). Status: FCM_FridgeINFO (32762), FCM_Fridge6 (32313). Line numbers refer to the reviewed generated C.'],
    ['Future devices', 'Keep these column names and add a Device / Variant column when combining devices. Shared IDs need a device/layout discriminator. Add confirmed capture references and revision history as evidence becomes available.'],
    ['Regenerate', 'Install openpyxl in your chosen Python environment, then run python3 scripts/export-lin-fridge-example.py. The generator reads the catalogue for byte descriptions; detailed fields are curated in the script.'],
]
sheet('Read Me', 'LIN TABLE LAYOUT | Dometic refrigerator',
      'One-device proposal • EC600 PSU source-derived reference • Start with the Byte Map tab',
      ['Topic', 'Explanation'], guide, [29, 125])
contexts = [(11, 'Control', 'PSU → fridge'), (12, 'Status', 'Fridge → PSU'), (11, 'Status echo', 'PSU → fridge')]
frames = []
for ident, context, direction in contexts:
    frames.append([ident, f'0x{ident:02X}', '0x8B' if ident == 11 else '0x4C', context, direction,
                   8, 'Enhanced', 'PSU',
                   {'Control': 'Settings change pending; uses control byte layout.',
                    'Status': 'Fridge supplies information in response to the PSU header.',
                    'Status echo': 'No change pending; PSU echoes ID 12 data with status meanings.'}[context],
                   'Not established', entries[ident]['layouts'][0]['source']])
sheet('Frames', 'FRAME SUMMARY | Control, status and echo',
      'ID and PID are shown separately. Payload length excludes checksum. Direction refers to response data, not the header.',
      ['LIN ID dec', 'LIN ID hex', 'Expected PID', 'Context', 'Data direction', 'Payload bytes', 'Checksum', 'Header sender', 'When / interpretation', 'Period (ms)', 'Source routines / lines'],
      frames, [12, 13, 14, 16, 20, 12, 14, 14, 54, 20, 46])
byte_rows = []
for ident, context, direction in contexts:
    mapping = entries[12 if context == 'Status echo' else ident]['layouts'][0]
    for b, meaning in enumerate(mapping['bytes']):
        # These rows already have an explicit context; avoid cross-context wording.
        if context == 'Control':
            meaning = {3: 'Written as 0. Old hour-setting code disabled.',
                       4: 'Written as 0. Old minute-setting code disabled.',
                       6: 'Written as 0 (spare) in control frame.',
                       7: 'Written as 0 (spare) in control frame.'}.get(b, meaning)
        byte_rows.append([ident, context, direction, f'B{b}', meaning,
                          'Source-derived; some meanings unknown' if b >= 5 and context != 'Control' else 'Source-derived', mapping['source']])
sheet('Byte Map', 'BYTE MAP | Familiar CAN-style view',
      'One row per byte for each context. Filter Context to see Control, Status or Status echo. The Signals tab splits packed bytes into fields.',
      ['LIN ID dec', 'Context', 'Data direction', 'Byte', 'Function / values', 'Evidence', 'Source routines / lines'],
      byte_rows, [12, 17, 20, 9, 85, 29, 48])

# byte, bits, mask, signal, extraction, options/conversion, evidence
mode = [
    (0, 'b2..0', '0x07', 'Operating mode', 'B0 & 0x07', '0=off; 1=auto; 3=gas; 5=12 V; 7=230 V. Other codes: not established.', 'Source-derived'),
    (0, 'b0', '0x01', 'On flag (overlaps mode)', 'B0 & 0x01', '0=off; 1=on. This bit is part of the mode field, not a separate setting.', 'Source-derived'),
    (0, 'b7..3', '0xF8', 'Upper mode-byte bits', '(B0 & 0xF8) >> 3', 'Unknown; do not assign meanings.', 'Unknown'),
]
control = mode + [
    (1, 'b7..0', '0xFF', 'Cooling setting', 'B1', 'Sent value = SetOut4[1] + 1. Allowed range not established.', 'Source-derived'),
    (2, 'b7..0', '0xFF', 'Frame-heater setting', 'B2', 'SetOut4[2] for large fridge/freezer; otherwise 0. Setting codes not established.', 'Source-derived; codes unknown'),
    (3, 'b7..0', '0xFF', 'Control constant', 'B3', 'Written as 0; old hour-setting code disabled.', 'Source-derived'),
    (4, 'b7..0', '0xFF', 'Control constant', 'B4', 'Written as 0; old minute-setting code disabled.', 'Source-derived'),
    (5, 'b4', '0x10', 'Automatic mode available', '(B5 & 0x10) >> 4', '1=available. This control path sets this bit in both 0x10 and 0x14.', 'Source-derived'),
    (5, 'b2', '0x04', 'Frame heater fitted', '(B5 & 0x04) >> 2', '0=not fitted; 1=fitted. Whole byte is 0x10 or 0x14.', 'Source-derived'),
    (5, 'b7..5, b3, b1..0', '0xEB', 'Other control bits', 'B5 & 0xEB', 'Written as 0 by this control path; no broader meaning established.', 'Source-derived'),
    (6, 'b7..0', '0xFF', 'Spare in control', 'B6', 'Written as 0.', 'Source-derived'),
    (7, 'b7..0', '0xFF', 'Spare in control', 'B7', 'Written as 0.', 'Source-derived'),
]
status = mode + [
    (1, 'b7..0', '0xFF', 'Cooling level', 'B1', 'PSU setting = B1 - 1; B1=0 is treated as starting up. Allowed range not established.', 'Source-derived'),
    (2, 'b7..0', '0xFF', 'Frame-heater setting', 'B2', 'Copied to SetOut4[2]. Setting codes not established.', 'Source-derived; codes unknown'),
    (3, 'b7..0', '0xFF', 'Fridge error code', 'B3', 'Error enumeration not established. PSU ignores errors during the first 15 seconds after on.', 'Source-derived; codes unknown'),
    (4, 'b2..0', '0x07', 'FridgeState', 'B4 & 0x07', 'Raw state code 0..7; meanings not established.', 'Source-derived; codes unknown'),
    (4, 'b3', '0x08', 'Manual mode flag', '(B4 & 0x08) >> 3', '0=flag clear; 1=manual mode flag set.', 'Source-derived'),
    (4, 'b4', '0x10', 'Door open flag', '(B4 & 0x10) >> 4', '0=flag clear; 1=door open flag set.', 'Source-derived'),
    (4, 'b7..5', '0xE0', 'Other status bits', '(B4 & 0xE0) >> 5', 'Unknown; not decoded by the reviewed PSU code.', 'Unknown'),
] + [(b, 'b7..0', '0xFF', 'Stored / echoed byte', f'B{b}', 'Unknown response-field meaning; stored and echoed. Do not apply control B5 flags.', 'Unknown') for b in (5, 6, 7)]
signal_rows = []
for ident, context, direction in contexts:
    source = entries[11 if context == 'Control' else 12]['layouts'][0]['source']
    for b, bits, mask, name, extract, options, evidence in control if context == 'Control' else status:
        signal_rows.append([ident, context, f'B{b}', bits, mask, name, extract, options, evidence, source])
sheet('Signals', 'SIGNAL DICTIONARY | Bits, masks and values',
      'One row per field. The On flag overlaps Operating mode intentionally. Expressions use bitwise & and right shift >>. Unestablished values remain explicit.',
      ['LIN ID dec', 'Context', 'Byte', 'Bits', 'Mask hex', 'Signal / variable', 'Extract raw value', 'Options / conversion / notes', 'Evidence', 'Source routines / lines'],
      signal_rows, [12, 16, 9, 24, 12, 30, 25, 74, 28, 46])

examples = [
    ['A', 11, '0x8B', 'Control', 'PSU → fridge', '03 03 00 00 00 10 00 00', 'B0=0x03 → gas; on=1. B1=3 → PSU cooling setting 2. B2=0 → raw frame-heater setting. B5=0x10 → auto available=1, frame heater fitted=0. B3/B4/B6/B7=0 control constants.'],
    ['B', 12, '0x4C', 'Status', 'Fridge → PSU', '03 03 00 00 18 00 00 00', 'B0=0x03 → gas; on=1. B1=3 → PSU cooling setting 2. B3=0 → raw error code (meaning not established). B4=0x18 = binary 00011000 → door flag=1, manual flag=1, FridgeState=0 (meaning unknown). B5..B7=0 have unknown meanings.'],
    ['C', 11, '0x8B', 'Status echo', 'PSU → fridge', '03 03 00 00 18 00 00 00', 'Same data as B, now echoed by the PSU on ID 11. Decode with the Status layout: B4 still contains status flags. B5=0 is unknown status data, not a declaration that auto mode is unavailable.'],
]
sheet('Examples', 'WORKED EXAMPLES | Synthetic payloads',
      'Illustrations only, not recorded frames. B0 is leftmost. Eight payload bytes shown; PID, checksum and receive validity must be checked separately in captures.',
      ['Example', 'LIN ID dec', 'Expected PID', 'Context', 'Data direction', 'B0 B1 B2 B3 B4 B5 B6 B7 (hex)', 'Decoded explanation'],
      examples, [12, 12, 14, 16, 20, 38, 110])
wb.active = 2  # Open on the familiar one-row-per-byte view.
wb.save(OUTPUT)

# Reopen the actual deliverable and check structure, contexts and byte coverage.
check = load_workbook(OUTPUT)
assert check.sheetnames == ['Read Me', 'Frames', 'Byte Map', 'Signals', 'Examples']
assert check['Byte Map'].max_row == 28
assert len(check['Frames'].tables) == 1
for ident, context, _ in contexts:
    rows = [r for r in check['Byte Map'].iter_rows(min_row=5, values_only=True) if r[:2] == (ident, context)]
    assert [r[3] for r in rows] == [f'B{b}' for b in range(8)]
for ws in check:
    assert ws.freeze_panes and len(ws.tables) == 1
for example in examples:
    assert len(bytes.fromhex(example[5])) == 8
assert (0x18 & 0x10) >> 4 == 1 and (0x18 & 0x08) >> 3 == 1
print(f'Created and reopened {OUTPUT}: {len(byte_rows)} byte rows, {len(signal_rows)} signal rows, {len(examples)} worked examples.')
